# -*- coding: utf-8 -*-
"""文档 / 数据附件解析。

支持：docx / pdf / txt / md / csv 文本抽取，以及 xlsx 的紧凑结构描述；
并提供「竞赛题目目录」扫描：每个题目 = 一份 PDF(题目) + 附件/(数据/模板)。
"""
from __future__ import annotations

import base64
from pathlib import Path
from typing import Optional

SUPPORTED_EXTS = {".docx", ".pdf", ".txt", ".md", ".csv", ".xlsx"}
ATTACHMENT_EXTS = {".xlsx", ".docx", ".doc", ".pdf", ".txt", ".md", ".csv"}


def extract_text(path) -> str:
    """按扩展名抽取文档/表格内容（xlsx 返回结构描述而非全量数据）。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        return _describe_xlsx(p)
    if suffix == ".docx":
        return _extract_docx(p)
    if suffix == ".pdf":
        return _extract_pdf(p)
    if suffix in (".txt", ".md", ".csv"):
        return p.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"不支持的文件类型: {suffix}（支持: {', '.join(sorted(SUPPORTED_EXTS))}）")


# ---------------------------------------------------------------------------
# docx / pdf
# ---------------------------------------------------------------------------

def _para_text(p_el) -> str:
    buf: list[str] = []
    for node in p_el.iter():
        tag = node.tag.split("}")[-1]
        if tag == "t":
            buf.append(node.text or "")
        elif tag == "br":
            buf.append("\n")
        elif tag == "tab":
            buf.append("\t")
    return "".join(buf)


def _extract_docx(p: Path) -> str:
    try:
        import docx
    except ImportError:
        raise RuntimeError("缺少 python-docx，请先执行 pip install -r requirements.txt")
    doc = docx.Document(str(p))
    parts: list[str] = []
    for el in doc.element.body.iter():
        if el.tag.endswith("}p"):
            text = _para_text(el).strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_pdf(p: Path) -> str:
    """抽取 PDF 文本（优先 pymupdf，回退 pypdf）。"""
    try:
        import pymupdf as fitz
        doc = fitz.open(str(p))
        parts = [page.get_text() for page in doc]
        doc.close()
        return "\n".join(parts)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("缺少 pymupdf（或 pypdf），请先执行 pip install -r requirements.txt")
    reader = PdfReader(str(p))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    return "\n".join(parts)


def extract_problem_pdf(pdf_path) -> dict:
    """抽取题目 PDF，返回 {'text': 全文, 'page_images': [每页 PNG 的 base64]}。

    整份 PDF 全部渲染成图，供视觉模型读取任意位置的公式（分式/上下标/幂次
    在文本层会被拉平、在 OCR 下易认错，直接看图最准）。
    """
    try:
        import pymupdf as fitz
    except ImportError:
        return {"text": _extract_pdf(Path(pdf_path)), "page_images": []}

    doc = fitz.open(str(pdf_path))
    texts: list[str] = []
    images: list[str] = []
    for page in doc:
        texts.append(page.get_text())
        pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
        images.append(base64.b64encode(pix.tobytes("png")).decode("ascii"))
    doc.close()
    return {"text": "\n".join(texts), "page_images": images}


# ---------------------------------------------------------------------------
# xlsx 结构描述（紧凑，供 LLM 理解数据而不灌全量）
# ---------------------------------------------------------------------------

def _fmt_row(vals, limit: int = 10) -> str:
    out: list[str] = []
    for i, v in enumerate(vals):
        if i >= limit:
            out.append(f"…（共{len(vals)}列）")
            break
        s = "" if v is None else str(v)
        if len(s) > 24:
            s = s[:24] + "…"
        out.append(s)
    return " | ".join(out)


def _describe_xlsx(p: Path, sample_rows: int = 4, max_cols: int = 10) -> str:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("缺少 openpyxl，请先执行 pip install -r requirements.txt")

    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out: list[str] = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            try:
                nrows, ncols = ws.max_row or 0, ws.max_column or 0
            except Exception:
                nrows = ncols = 0
            out.append(f"· 工作表「{sheet_name}」：约 {nrows} 行 × {ncols} 列")
            try:
                rows = list(ws.iter_rows(min_row=1, max_row=min(nrows, sample_rows + 1), values_only=True))
            except Exception:
                rows = []
            if rows:
                out.append("  列头: " + _fmt_row(rows[0] or [], max_cols))
            non_empty = 0
            for r in rows[1:sample_rows + 1]:
                vals = r or []
                non_empty += sum(1 for c in vals if c not in (None, ""))
                out.append("  样例: " + _fmt_row(vals, max_cols))
            if non_empty == 0:
                out.append("  （数据区为空，可能是待填充的结果模板）")
    finally:
        wb.close()
    return "\n".join(out)


# ---------------------------------------------------------------------------
# 竞赛题目目录扫描
# ---------------------------------------------------------------------------

def scan_problem_dir(root: Path) -> list[dict]:
    """扫描 竞赛题目/ 目录，返回每个题目的信息。

    目录约定：
        竞赛题目/
          A题/A题.pdf + 附件/*.xlsx|.docx|...
          B题/B题.pdf + 附件/*.docx|...
        （也兼容直接把题目 PDF 放在 root 下、附件放在同级 附件/ 的情况）

    每个题目返回 {"title": 目录名, "problem_text": 题目正文, "attachments": [{"name","desc"}]}
    """
    problems: list[dict] = []
    if not root.is_dir():
        return problems

    # 1) 一级子目录各算一题
    for entry in sorted(root.iterdir()):
        if entry.is_dir():
            p = _scan_one_problem(entry)
            if p:
                problems.append(p)

    # 2) root 下直接放题目 PDF 的情况
    for entry in sorted(root.iterdir()):
        if entry.is_file() and entry.suffix.lower() == ".pdf":
            p = _scan_one_problem(root, problem_pdf=entry)
            if p:
                problems.append(p)

    return problems


def _scan_one_problem(problem_dir: Path, problem_pdf: Optional[Path] = None) -> Optional[dict]:
    title = problem_dir.name

    if problem_pdf is None:
        # 优先 <目录名>.pdf，否则取目录直属第一个 pdf
        candidate = problem_dir / f"{title}.pdf"
        if not candidate.exists():
            direct = [p for p in sorted(problem_dir.glob("*.pdf"))]
            if not direct:
                return None
            candidate = direct[0]
        problem_pdf = candidate

    try:
        extracted = extract_problem_pdf(problem_pdf)
    except Exception:
        return None
    problem_text = extracted.get("text", "")
    page_images = extracted.get("page_images", [])
    if not problem_text.strip():
        return None

    attachments = []
    att_dir = problem_dir / "附件"
    if att_dir.is_dir():
        for f in sorted(att_dir.rglob("*")):
            if not f.is_file() or f.suffix.lower() not in ATTACHMENT_EXTS:
                continue
            try:
                desc = extract_text(f)
            except Exception as exc:
                desc = f"(failed to inspect attachment: {exc})"
            attachments.append({
                "name": str(f.relative_to(problem_dir)),
                "path": str(f.resolve()),
                "suffix": f.suffix.lower(),
                "size": f.stat().st_size,
                "desc": desc,
            })

    return {
        "title": title,
        "problem_text": problem_text,
        "attachments": attachments,
        "page_images": page_images,
        "problem_pdf": str(Path(problem_pdf).resolve()),
    }

def build_problem_input(problem: dict) -> str:
    """把题目信息拼接为送给 LLM 的完整输入文本。"""
    parts = [f"# 题目：{problem.get('title', '')}", "", (problem.get("problem_text") or "").strip()]
    atts = problem.get("attachments") or []
    if atts:
        lines = ["", "## 附件文件内容（独立数据文件 / 结果模板，已读取其结构）"]
        for a in atts:
            lines.append(f"\n### {a['name']}\n{a['desc']}")
        parts.append("\n".join(lines))
    return "\n".join(parts)
