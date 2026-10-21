# -*- coding: utf-8 -*-
"""
2026 年全国大学生数学建模竞赛 D 题 —— 时频冲突检测与消解
解题主程序（纯标准库 + openpyxl，单文件、单入口）

题目要点
    频域：把有限带宽离散为 100 个宽度 Δf 的频段；一个用频计划占用连续若干频段，且占用频段数不变；
    时域：以 Δt 为最小时间单位；一个用频计划由「频段区间 / 首次使用时间区间 / 间隔时长 / 使用次数」四参数描述；
    冲突：两个用频计划在频段区间与时间区间上同时存在交叠；第 k 次使用的时间区间为
          [t0 + k·iv, t0 + k·iv + tau)，其中 (t0, tau, iv) 分别为首次使用起点、单次时长、间隔时长；
    消解：调整参数或撤销计划，且一个计划只能调整其中一个参数或撤销；频段 / 时间最大平移
          幅度分别为 10Δf、5Δt；一般不调整使用次数与间隔时长；
          优先级 A > B > C，目标是「少撤销 → 少调整 → 幅度小 → 高优先级尽量不动」。

四问
    问题 1  时频冲突检测                                   -> result1.xlsx
    问题 2  幅度受限下的冲突消解（频移 / 时移 / 撤销）        -> result2.xlsx
    问题 3  在问题 2 无冲突计划上做 C 类容量最大化（不限幅度） -> result3.xlsx
    问题 4  允许 C 类调整间隔时长（|Δiv| ≤ 10Δt）的消解      -> result4.xlsx

数据
    附件1.xlsx：150 件装备（A 类 20、B 类 40、C 类 90）
    字段：用频装备编号 / 频段区间 / 时间区间 / 间隔时长 / 使用次数

运行
    python main.py                      # 自动在若干候选路径中查找 附件1.xlsx
    python main.py 路径\附件1.xlsx       # 手动指定附件路径

输出
    result1.xlsx ~ result4.xlsx（与本脚本同目录），并在控制台打印四问的统计报告。
"""

import os
import re
import sys
import bisect
from copy import deepcopy

import openpyxl

try:                                # 保证中文 / 希腊字母在控制台正常输出
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass


# ============================================================== 0. 路径与常量
HERE = os.path.dirname(os.path.abspath(__file__))
INPUT_CANDIDATES = [
    os.path.join(HERE, '附件1.xlsx'),
    os.path.join(HERE, '附件', '附件1.xlsx'),
    os.path.join(HERE, '..', '附件', '附件1.xlsx'),
    os.path.join(HERE, '..', '..', '附件1.xlsx'),
    os.path.join(HERE, '..', '..', '2026D题演示', '附件1.xlsx'),
    r'C:\tidiwork\problems\2026\D题\附件\附件1.xlsx',
]
OUT_DIR = HERE

FREQ_SLOTS = 100        # 频域可用频段数
MAX_FREQ_SHIFT = 10     # 频段最大平移幅度（单位 Δf）
MAX_TIME_SHIFT = 5      # 时间最大平移幅度（单位 Δt）
MAX_IV_SHIFT = 10       # 问题四允许的间隔时长最大调整量（单位 Δt）
PRIORITY = {'A': 3, 'B': 2, 'C': 1}


# ============================================================== 1. 读取附件
def find_input():
    for p in INPUT_CANDIDATES:
        if os.path.exists(p):
            return p
    raise FileNotFoundError('未找到 附件1.xlsx，请用 python main.py <附件路径> 指定')


def _parse_interval(text):
    """把 '[80,90)' / '[80, 90]' 之类的区间字符串解析为 (80, 90)。"""
    nums = re.findall(r'-?\d+', str(text))
    if len(nums) < 2:
        raise ValueError('无法解析区间：%r' % (text,))
    return int(nums[0]), int(nums[1])


def num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    return int(float(str(v).strip()))


def load_data(path):
    """读取附件 1，返回用频计划列表。

    每个计划为 dict：id / cls / priority / fs / fe（频段区间）/ ts / te（首次使用时间区间）
    / iv（间隔时长）/ n（使用次数）/ status / adj（已调整的参数类型）。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    plans = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or row[0] is None:
            continue
        pid = str(row[0]).strip()
        fs, fe = _parse_interval(row[1])
        ts, te = _parse_interval(row[2])
        iv = num(row[3])
        n = num(row[4])
        cls = pid[0].upper()
        plans.append({
            'id': pid, 'cls': cls, 'priority': PRIORITY.get(cls, 0),
            'fs': fs, 'fe': fe, 'ts': ts, 'te': te, 'tau': te - ts,
            'iv': iv, 'n': n, 'status': 'active', 'adj': 'none',
        })
    plans.sort(key=lambda p: p['id'])
    return plans


# ============================================================== 2. 公共几何
def plan_rects(p):
    """把一个用频计划展开为若干时频矩形 (fs, fe, t0, t1)。"""
    out = []
    for k in range(p['n']):
        t0 = p['ts'] + k * p['iv']
        out.append((p['fs'], p['fe'], t0, t0 + p['tau']))
    return out


def rects_overlap(a, b):
    """两个时频矩形是否冲突（半开区间，频段与时间同时交叠才算冲突）。"""
    return (max(a[0], b[0]) < min(a[1], b[1])) and (max(a[2], b[2]) < min(a[3], b[3]))


def plans_conflict(a, b):
    """两个用频计划是否存在时频冲突。"""
    if max(a['fs'], b['fs']) >= min(a['fe'], b['fe']):
        return False                      # 频段区间无交叠 -> 必不冲突
    for r1 in plan_rects(a):
        for r2 in plan_rects(b):
            if rects_overlap(r1, r2):
                return True
    return False


def find_all_conflicts(plans):
    """返回所有冲突计划对的 (i, j) 索引（i < j），按索引顺序排列。"""
    active = [p for p in plans if p['status'] == 'active']
    pairs = []
    for i in range(len(active)):
        for j in range(i + 1, len(active)):
            if plans_conflict(active[i], active[j]):
                pairs.append((i, j))
    return active, pairs


def is_conflict_free(plans):
    return not any(plans_conflict(plans[i], plans[j])
                   for i in range(len(plans)) for j in range(i + 1, len(plans)))


# ============================================================== 3. 问题一：冲突检测
def q1_detect(plans):
    active, pairs = find_all_conflicts(plans)
    rows = []
    for k, (i, j) in enumerate(pairs, 1):
        rows.append([k, active[i]['id'], active[j]['id']])

    # ---- 统计 ----
    stat = {'pairs': len(pairs), 'total': len(active), 'rects': sum(p['n'] for p in active)}
    hot = set()
    combo = {}
    cls_conf = {}
    for i, j in pairs:
        a, b = active[i], active[j]
        hot.add(a['id']); hot.add(b['id'])
        key = '-'.join(sorted([a['cls'], b['cls']]))
        combo[key] = combo.get(key, 0) + 1
        cls_conf.setdefault(a['cls'], 0)
    involved = {}
    for pid in hot:
        involved[pid[0]] = involved.get(pid[0], 0) + 1
    stat['involved'] = len(hot)
    stat['combo'] = combo
    stat['involved_cls'] = involved

    print('=' * 68)
    print('问题 1  时频冲突检测')
    print('=' * 68)
    print('  用频计划总数        : %d' % stat['total'])
    print('  时频矩形总数        : %d' % stat['rects'])
    print('  冲突计划对          : %d' % stat['pairs'])
    print('  涉及冲突的装备      : %d 件 (%.1f%%)'
          % (stat['involved'], 100.0 * stat['involved'] / stat['total']))
    print('  冲突对类别分布      : ' +
          '  '.join('%s×%d' % (k, v) for k, v in sorted(stat['combo'].items())))
    print('  各类涉及冲突装备数  : ' +
          '  '.join('%s类 %d 件' % (k, v) for k, v in sorted(stat['involved_cls'].items())))
    p1 = _write_xlsx('result1.xlsx',
                     ['序号', '冲突装备1', '冲突设备2'], rows)
    print('  -> 已写出 %s（%d 行）' % (os.path.basename(p1), len(rows)))
    return stat


# ============================================================== 4. 问题二 / 问题四：冲突消解
def _first_conflict(active):
    for i in range(len(active)):
        for j in range(i + 1, len(active)):
            if plans_conflict(active[i], active[j]):
                return i, j
    return None


def _try_adjust(p, others, allow_iv):
    """尝试平移 / 调整间隔消解 p 与其他计划的冲突，成功则就地修改并返回 True。

    调整顺序：频移（≤10Δf）→ 时移（≤5Δt）→ 间隔调（仅 C 类且 allow_iv，|Δiv|≤10Δt）。
    同一幅度下先试左移（-）再试右移（+），保证结果确定。
    """
    if p['adj'] != 'none':
        return False                       # 一个计划只能调整一个参数

    # (1) 频段平移
    for d in range(1, MAX_FREQ_SHIFT + 1):
        for s in (-1, 1):
            nfs, nfe = p['fs'] + s * d, p['fe'] + s * d
            if nfs < 0 or nfe > FREQ_SLOTS:
                continue
            fs0, fe0 = p['fs'], p['fe']
            p['fs'], p['fe'] = nfs, nfe
            if not any(plans_conflict(p, q) for q in others):
                p['adj'] = 'freq'
                return True
            p['fs'], p['fe'] = fs0, fe0

    # (2) 时间平移
    for d in range(1, MAX_TIME_SHIFT + 1):
        for s in (-1, 1):
            nts = p['ts'] + s * d
            if nts < 0:
                continue
            ts0 = p['ts']
            p['ts'] = nts
            if not any(plans_conflict(p, q) for q in others):
                p['adj'] = 'time'
                return True
            p['ts'] = ts0

    # (3) 间隔时长调整（仅 C 类，仅问题四）
    if allow_iv and p['cls'] == 'C':
        for d in range(1, MAX_IV_SHIFT + 1):
            for s in (-1, 1):
                niv = p['iv'] + s * d
                if niv < 1:
                    continue
                iv0 = p['iv']
                p['iv'] = niv
                if not any(plans_conflict(p, q) for q in others):
                    p['adj'] = 'iv'
                    return True
                p['iv'] = iv0
    return False


def resolve(plans, allow_iv=False):
    """贪心消解：反复取首个冲突对，让低优先级一方让路，直至无冲突。

    每次迭代：优先调整尚未调整过的低优先级计划；若其无法调整则改用另一方；
    两方都不可调整时撤销低优先级计划。每个计划至多被调整一次或被撤销。
    """
    active = [p for p in plans if p['status'] == 'active']
    guard = 0
    while True:
        guard += 1
        if guard > 10 * len(plans):
            raise RuntimeError('消解迭代次数异常')
        pair = _first_conflict(active)
        if pair is None:
            break
        i, j = pair
        a, b = active[i], active[j]
        # 候选让路者：优先级低者优先；优先级相同取序号靠后者
        if a['priority'] != b['priority']:
            cands = [a, b] if a['priority'] < b['priority'] else [b, a]
        else:
            cands = [b, a]
        solved = False
        for loser in cands:
            others = [q for q in active if q is not loser and q['status'] == 'active']
            if _try_adjust(loser, others, allow_iv):
                solved = True
                break
        if not solved:
            cands[0]['status'] = 'cancelled'
            active = [p for p in active if p['status'] == 'active']
    return plans


def _shift_rows(plans, with_iv=False):
    """整理 result2 / result4 的行：仅列出被调整或被撤销的计划。"""
    rows = []
    for p in plans:
        if p['adj'] == 'none' and p['status'] == 'active':
            continue
        freq = '[%d,%d)' % (p['fs'], p['fe']) if p['adj'] == 'freq' else ''
        time = '[%d,%d)' % (p['ts'], p['ts'] + p['tau']) if p['adj'] == 'time' else ''
        row = [p['id'], freq, time]
        if with_iv:
            row.append(p['iv'] if p['adj'] == 'iv' else '')
        row.append('是' if p['status'] == 'cancelled' else '')
        rows.append(row)
    return rows


def _stat_table(plans, title, rows_path):
    print('  %s' % title)
    print('    %-6s %-8s %-8s %-8s' % ('类别', '保留数量', '调整数量', '撤销数量'))
    total = {'keep': 0, 'adj': 0, 'cancel': 0}
    for c in ('A', 'B', 'C'):
        keep = sum(1 for p in plans if p['cls'] == c and p['status'] == 'active' and p['adj'] == 'none')
        adj = sum(1 for p in plans if p['cls'] == c and p['status'] == 'active' and p['adj'] != 'none')
        can = sum(1 for p in plans if p['cls'] == c and p['status'] == 'cancelled')
        total['keep'] += keep; total['adj'] += adj; total['cancel'] += can
        print('    %-6s %-10d %-10d %-10d' % (c + ' 类', keep, adj, can))
    print('    %-6s %-10d %-10d %-10d' % ('合计', total['keep'], total['adj'], total['cancel']))
    return total


def q2_resolve(plans):
    data = deepcopy(plans)
    resolve(data, allow_iv=False)
    assert is_conflict_free([p for p in data if p['status'] == 'active']), '问题 2 仍有冲突'

    print('=' * 68)
    print('问题 2  幅度受限下的冲突消解（频移 ≤10Δf，时移 ≤5Δt）')
    print('=' * 68)
    rows = _shift_rows(data, with_iv=False)
    total = _stat_table(data, '时频冲突消解的统计结果（表 1 格式）', None)
    adj_kind = {}
    for p in data:
        if p['status'] == 'active' and p['adj'] != 'none':
            adj_kind[p['adj']] = adj_kind.get(p['adj'], 0) + 1
    print('    调整方式：频移 %d 件，时移 %d 件'
          % (adj_kind.get('freq', 0), adj_kind.get('time', 0)))
    p2 = _write_xlsx('result2.xlsx',
                     ['用频装备编号', '调整后频段区间', '调整后时间区间', '是否撤销用频计划'], rows)
    print('  -> 已写出 %s（%d 行）' % (os.path.basename(p2), len(rows)))
    return data, total


# ============================================================== 5. 问题三：容量最大化
def q3_capacity(data):
    active = [p for p in data if p['status'] == 'active']
    max_t = max((r[3] for p in active for r in plan_rects(p)), default=0)
    total_area = FREQ_SLOTS * max_t
    used_area = sum((r[1] - r[0]) * (r[3] - r[2]) for p in active for r in plan_rects(p))
    free_area = max(0, total_area - used_area)

    # C 类模板：附件 1 中 C 类装备参数完全一致（频宽 3、时长 2、间隔 8、使用 12 次）
    cset = [p for p in active if p['cls'] == 'C']
    tpl = {'bw': cset[0]['fe'] - cset[0]['fs'], 'tau': cset[0]['tau'],
           'iv': cset[0]['iv'], 'n': cset[0]['n']}
    span = tpl['iv'] * (tpl['n'] - 1) + tpl['tau']        # 单件占用的时间跨度
    per = tpl['bw'] * tpl['tau'] * tpl['n']               # 单件占用的时频面积

    # 占用表：freq -> 已占用的时间区间（互不重叠且有序）
    occ = {f: [] for f in range(FREQ_SLOTS)}
    for p in active:
        for r in plan_rects(p):
            for f in range(r[0], r[1]):
                bisect.insort(occ[f], (r[2], r[3]))

    def fits(fs, ts):
        t_end = ts + span
        if t_end > max_t:
            return False
        for k in range(tpl['n']):
            t0 = ts + k * tpl['iv']
            t1 = t0 + tpl['tau']
            for f in range(fs, fs + tpl['bw']):
                L = occ[f]
                pos = bisect.bisect_left(L, (t0,))
                for idx in (pos - 1, pos):
                    if 0 <= idx < len(L):
                        a, b = L[idx]
                        if max(a, t0) < min(b, t1):
                            return False
        return True

    def place(fs, ts):
        for k in range(tpl['n']):
            t0 = ts + k * tpl['iv']
            for f in range(fs, fs + tpl['bw']):
                bisect.insort(occ[f], (t0, t0 + tpl['tau']))

    placed = []
    for fs in range(0, FREQ_SLOTS - tpl['bw'] + 1, tpl['bw']):
        for ts in range(0, max(0, max_t - span) + 1):
            if fits(fs, ts):
                place(fs, ts)
                placed.append((fs, fs + tpl['bw'], ts, ts + tpl['tau'], ts + span))

    feasible = len(placed)
    ub = free_area // per
    density = (used_area + feasible * per) / float(total_area) * 100 if total_area else 0.0

    print('=' * 68)
    print('问题 3  无冲突计划上的 C 类容量最大化（不限平移幅度）')
    print('=' * 68)
    print('  时频资源总量        : 100 Δf × %d Δt = %d 个单元' % (max_t, total_area))
    print('  已占用（问题 2 后）  : %d 个单元（占用率 %.1f%%）' % (used_area, 100.0 * used_area / total_area))
    print('  剩余可分配空间      : %d 个单元' % free_area)
    print('  C 类单件模板        : 频宽 %dΔf × 时长 %dΔt × %d 次（面积 %d，跨度 %dΔt）'
          % (tpl['bw'], tpl['tau'], tpl['n'], per, span))
    print('  容量上界（面积比）  : %d 件' % ub)
    print('  贪心可行解          : %d 件（装填密度 %.1f%%，与上界差 +%d 件）'
          % (feasible, density, ub - feasible))
    rows = [[k, '[%d,%d)' % (fs, fe), '[%d,%d)' % (ts, te)]
            for k, (fs, fe, ts, te, _) in enumerate(placed, 1)]
    p3 = _write_xlsx('result3.xlsx',
                     ['新增用频装备序号', '调整后频段区间', '调整后时间区间'], rows)
    print('  -> 已写出 %s（%d 行）' % (os.path.basename(p3), len(rows)))
    return {'feasible': feasible, 'ub': ub, 'density': density, 'free': free_area,
            'used': used_area, 'total': total_area, 'max_t': max_t, 'per': per}


# ============================================================== 6. 问题四：间隔调整
def q4_interval(plans):
    data = deepcopy(plans)
    resolve(data, allow_iv=True)
    assert is_conflict_free([p for p in data if p['status'] == 'active']), '问题 4 仍有冲突'

    print('=' * 68)
    print('问题 4  允许 C 类调整间隔时长（|Δiv| ≤ 10Δt）的消解')
    print('=' * 68)
    rows = _shift_rows(data, with_iv=True)
    total = _stat_table(data, '时频冲突消解的统计结果（表 1 格式）', None)
    adj_kind = {}
    for p in data:
        if p['status'] == 'active' and p['adj'] != 'none':
            adj_kind[p['adj']] = adj_kind.get(p['adj'], 0) + 1
    print('    调整方式：频移 %d 件，时移 %d 件，间隔调整 %d 件'
          % (adj_kind.get('freq', 0), adj_kind.get('time', 0), adj_kind.get('iv', 0)))
    p4 = _write_xlsx('result4.xlsx',
                     ['用频装备编号', '调整后频段范围', '调整后时间区间', '调整后间隔时长', '是否撤销用频计划'],
                     rows)
    print('  -> 已写出 %s（%d 行）' % (os.path.basename(p4), len(rows)))
    return data, total


# ============================================================== 7. 工具函数
def _write_xlsx(name, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(header)
    for r in rows:
        ws.append(r)
    path = os.path.join(OUT_DIR, name)
    wb.save(path)
    return path


# ============================================================== 8. 入口
def main():
    path = sys.argv[1] if len(sys.argv) > 1 else find_input()
    print('附件来源：%s' % path)
    plans = load_data(path)
    print('读入用频计划 %d 条（A %d / B %d / C %d）'
          % (len(plans),
             sum(1 for p in plans if p['cls'] == 'A'),
             sum(1 for p in plans if p['cls'] == 'B'),
             sum(1 for p in plans if p['cls'] == 'C')))

    q1_detect(plans)
    p2_data, p2_stat = q2_resolve(plans)
    q3_capacity(p2_data)
    q4_interval(plans)

    print('=' * 68)
    print('全部完成：result1.xlsx / result2.xlsx / result3.xlsx / result4.xlsx')
    print('已写入目录 %s' % OUT_DIR)


if __name__ == '__main__':
    main()
