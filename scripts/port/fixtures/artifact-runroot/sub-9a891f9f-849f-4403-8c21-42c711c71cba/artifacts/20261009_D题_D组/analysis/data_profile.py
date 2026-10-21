# -*- coding: utf-8 -*-
"""S1 输入盘点用的**只读**数据画像脚本（不是解题代码：只做字段/取值/覆盖统计，不做冲突检测、不做优化、不求解）。

用途：为题目分析提供附件1 的字段覆盖率、类别分布、区间解析可行性、带宽/时长分布，
以及由数据推得的频率轴与时间轴占用范围，供 problem_contract 与数据质量评估引用。

用法：
    run.cmd <本脚本绝对路径> --xlsx <附件1绝对路径> --out <data_profile.json 绝对路径>
退出码：0 成功；1 读取失败
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

INTERVAL = re.compile(r"^\s*[\[(]\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)\s*[\])]\s*$")


def _fix(stream):
    if stream and hasattr(stream, "reconfigure"):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def parse_interval(raw):
    if raw is None:
        return None
    m = INTERVAL.match(str(raw))
    if not m:
        return None
    return float(m.group(1)), float(m.group(2))


def cls_of(rid):
    m = re.match(r"[A-Za-z]+", rid or "")
    return m.group(0) if m else "?"


def main() -> int:
    _fix(sys.stdout)
    _fix(sys.stderr)
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError as exc:
        print("! 缺少 openpyxl：", exc)
        return 1

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    names = list(wb.sheetnames)
    ws = wb[names[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        print("! 空表")
        return 1

    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: i for i, h in enumerate(headers)}
    body = [r for r in rows[1:] if any(v is not None and str(v).strip() != "" for v in r)]

    def cell(r, key):
        i = idx.get(key)
        return r[i] if i is not None and len(r) > i else None

    ids = [str(cell(r, "用频装备编号") or "").strip() for r in body]
    dup = [k for k, v in Counter(ids).items() if v > 1]
    class_counts = Counter(cls_of(i) for i in ids)

    band_bad, time_bad, gap_bad, use_bad = [], [], [], []
    band_w, time_d = Counter(), Counter()
    gap_vals, use_vals = Counter(), Counter()
    bmin = bmax = tmin = tmax = None
    max_last_end = None
    per_class = {}
    tf_cells_total = 0.0
    band_out_of_range = []
    time_nonpositive = []
    gap_lt_duration = []
    rec_count = 0

    for rid, r in zip(ids, body):
        b = parse_interval(cell(r, "频段区间"))
        t = parse_interval(cell(r, "时间区间"))
        if b is None:
            band_bad.append((rid, str(cell(r, "频段区间"))))
        if t is None:
            time_bad.append((rid, str(cell(r, "时间区间"))))
        g_raw, u_raw = cell(r, "间隔时长"), cell(r, "使用次数")
        try:
            g = float(g_raw)
        except (TypeError, ValueError):
            g = None
            gap_bad.append((rid, str(g_raw)))
        try:
            u = float(u_raw)
        except (TypeError, ValueError):
            u = None
            use_bad.append((rid, str(u_raw)))
        if g is not None:
            gap_vals[g] += 1
        if u is not None:
            use_vals[u] += 1
        if b is None or t is None or g is None or u is None:
            continue
        rec_count += 1
        bw, td = b[1] - b[0], t[1] - t[0]
        band_w[bw] += 1
        time_d[td] += 1
        bmin = b[0] if bmin is None else min(bmin, b[0])
        bmax = b[1] if bmax is None else max(bmax, b[1])
        tmin = t[0] if tmin is None else min(tmin, t[0])
        tmax = t[1] if tmax is None else max(tmax, t[1])
        if bw <= 0 or td <= 0:
            time_nonpositive.append((rid, b, t))
        # 附录口径：间隔时长 g 是相邻两次使用之间的空闲时长，故周期 = d + g
        # 第 k 次占用窗口 = [t0 + (k-1)(d+g), t0 + (k-1)(d+g) + d)，k = 1..u
        last_end = t[1] + (u - 1) * (g + td)
        max_last_end = last_end if max_last_end is None else max(max_last_end, last_end)
        if g < td:
            gap_lt_duration.append((rid, g, td))
        tf_cells_total += bw * td * u
        if b[1] > 100 or b[0] < 0:
            band_out_of_range.append((rid, b))
        c = per_class.setdefault(cls_of(rid), {
            "n": 0, "band_width": Counter(), "first_duration": Counter(),
            "gap": Counter(), "use_count": Counter(),
            "band_min": None, "band_max": None,
            "first_t_min": None, "first_t_max": None, "max_last_end": None,
        })
        c["n"] += 1
        c["band_width"][bw] += 1
        c["first_duration"][td] += 1
        c["gap"][g] += 1
        c["use_count"][u] += 1
        c["band_min"] = b[0] if c["band_min"] is None else min(c["band_min"], b[0])
        c["band_max"] = b[1] if c["band_max"] is None else max(c["band_max"], b[1])
        c["first_t_min"] = t[0] if c["first_t_min"] is None else min(c["first_t_min"], t[0])
        c["first_t_max"] = t[1] if c["first_t_max"] is None else max(c["first_t_max"], t[1])
        c["max_last_occupancy_end"] = last_end if c["max_last_end"] is None else max(c["max_last_end"], last_end)
        c["max_last_end"] = last_end if c["max_last_end"] is None else max(c["max_last_end"], last_end)

    def flat(counter):
        return {str(k): v for k, v in sorted(counter.items())}

    profile = {
        "purpose": "S1 输入盘点/数据质量画像（只读统计，非求解）",
        "source": Path(args.xlsx).name,
        "sheets": names,
        "headers": headers,
        "n_data_rows": len(body),
        "n_rows_fully_parsed": rec_count,
        "ids_first6": ids[:6],
        "ids_last6": ids[-6:],
        "duplicate_ids": dup,
        "class_prefix_counts": {k: v for k, v in sorted(class_counts.items())},
        "unparsed": {
            "band_interval": {"n": len(band_bad), "sample": band_bad[:5]},
            "time_interval": {"n": len(time_bad), "sample": time_bad[:5]},
            "gap": {"n": len(gap_bad), "sample": gap_bad[:5]},
            "use_count": {"n": len(use_bad), "sample": use_bad[:5]},
        },
        "band_axis_observed": {"min": bmin, "max": bmax},
        "first_time_axis_observed": {"min": tmin, "max": tmax},
        "derived_max_last_occupancy_end": max_last_end,
        "derived_sum_of_time_freq_cells": tf_cells_total,
        "band_width_distribution": flat(band_w),
        "first_time_duration_distribution": flat(time_d),
        "gap_distribution": flat(gap_vals),
        "use_count_distribution": flat(use_vals),
        "flags": {
            "band_end_gt_100_or_neg": band_out_of_range[:10],
            "band_or_time_nonpositive": time_nonpositive[:10],
            "gap_lt_first_duration": {"n": len(gap_lt_duration), "sample": gap_lt_duration[:10]},
        },
        "by_class": {
            k: {
                "n": v["n"],
                "band_width": flat(v["band_width"]),
                "first_duration": flat(v["first_duration"]),
                "gap": flat(v["gap"]),
                "use_count": flat(v["use_count"]),
                "band_min": v["band_min"], "band_max": v["band_max"],
                "first_t_min": v["first_t_min"], "first_t_max": v["first_t_max"],
                "max_last_occupancy_end": v["max_last_end"],
            }
            for k, v in sorted(per_class.items())
        },
    }
    Path(args.out).write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(profile, ensure_ascii=False, indent=2))
    print(f"写出：{Path(args.out)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
