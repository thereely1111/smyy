# -*- coding: utf-8 -*-
"""
Extract Word comments/annotations from a DOCX file.
Parses word/comments.xml and word/document.xml to get:
  - Comment text, author, date
  - The document context (text surrounding each comment anchor)

Usage:
  python extract_docx_comments.py <input.docx> [--output json|xlsx] [--output-path <path>]
"""
import sys
import os
import json
import zipfile
from lxml import etree

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _load_image_rel_map(zf):
    """读取 document.xml.rels，建立 rId → 图片文件名 的映射。"""
    rel_map = {}
    try:
        rels_xml = zf.read('word/_rels/document.xml.rels')
    except KeyError:
        return rel_map
    rels_root = etree.fromstring(rels_xml)
    for rel in rels_root:
        rtype = rel.get('Type', '')
        if 'image' in rtype:
            rel_map[rel.get('Id')] = rel.get('Target', '')
    return rel_map


def extract_comments(docx_path):
    """Extract all comments and their document context from a DOCX file.
    Returns a list of dicts: {id, author, date, text, context}.
    """
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"File not found: {docx_path}")

    with zipfile.ZipFile(docx_path, 'r') as zf:
        if 'word/comments.xml' not in zf.namelist():
            raise ValueError("No comments.xml found in the document. The document has no comments.")

        comments_xml = zf.read('word/comments.xml')
        doc_xml = zf.read('word/document.xml')
        image_rel_map = _load_image_rel_map(zf)

    # Parse comments
    comments_root = etree.fromstring(comments_xml)
    comments_info = {}
    for comment in comments_root.findall(f'{{{W_NS}}}comment'):
        cid = int(comment.get(f'{{{W_NS}}}id'))
        author = comment.get(f'{{{W_NS}}}author', '')
        date = comment.get(f'{{{W_NS}}}date', '')
        text_parts = []
        for t_elem in comment.iter(f'{{{W_NS}}}t'):
            if t_elem.text:
                text_parts.append(t_elem.text)
        comments_info[cid] = {
            'id': cid,
            'author': author,
            'date': date,
            'text': ''.join(text_parts),
        }

    # Parse document for comment references
    doc_root = etree.fromstring(doc_xml)
    paragraphs = []
    for para in doc_root.iter(f'{{{W_NS}}}p'):
        para_text = []
        comment_ids = []
        seen_range_starts = []
        image_rids = []
        for elem in para.iter():
            if elem.tag == f'{{{W_NS}}}commentRangeStart':
                cid_v = elem.get(f'{{{W_NS}}}id')
                if cid_v:
                    seen_range_starts.append(int(cid_v))
            elif elem.tag == f'{{{W_NS}}}commentReference':
                cid_v = elem.get(f'{{{W_NS}}}id')
                if cid_v:
                    comment_ids.append(int(cid_v))
            elif elem.tag == f'{{{W_NS}}}t' and elem.text:
                para_text.append(elem.text)
            elif elem.tag == f'{{{A_NS}}}blip':
                rid = elem.get(f'{{{R_NS}}}embed')
                if rid:
                    image_rids.append(rid)
        paragraphs.append({
            'text': ''.join(para_text).strip(),
            'comment_ids': list(set(comment_ids + seen_range_starts)),
            'image_rids': list(dict.fromkeys(image_rids)),
        })

    # Map comments to context
    comment_context = {}
    for i, pdata in enumerate(paragraphs):
        for cid in pdata['comment_ids']:
            if cid not in comment_context:
                comment_context[cid] = []
            comment_context[cid].append(i)

    # Build results
    results = []
    for cid in sorted(comments_info.keys()):
        ci = comments_info[cid]
        ctx_parts = []
        image_names = []
        anchor_para_index = None
        if cid in comment_context:
            for pi in comment_context[cid]:
                pdata = paragraphs[pi]
                text = pdata['text']
                # 图片锚点兜底：段落文本为空但有图片时，用图片文件名 + 前后段落文本补足 context
                if not text and pdata['image_rids']:
                    image_names = [image_rel_map.get(r, r) for r in pdata['image_rids']]
                    anchor_para_index = pi
                    prev_text = paragraphs[pi - 1]['text'] if pi - 1 >= 0 else ''
                    next_text = paragraphs[pi + 1]['text'] if pi + 1 < len(paragraphs) else ''
                    text = f"[图片:{'、'.join(image_names)}] {prev_text} {next_text}".strip()
                ctx_parts.append(text)
        context = ' '.join(ctx_parts)
        results.append({
            'id': cid,
            'author': ci['author'],
            'date': ci['date'],
            'text': ci['text'],
            'context': context,
            'image_names': image_names,
            'anchor_para_index': anchor_para_index,
        })

    return results


def output_json(results, path):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"JSON output: {path} ({len(results)} comments)")


def output_xlsx(results, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        print("openpyxl not installed. Install with: pip install openpyxl")
        output_json(results, path.replace('.xlsx', '.json'))
        return

    wb = Workbook()
    ws = wb.active
    ws.title = '批注汇总'

    # Header style
    hdr_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    hdr_fill = PatternFill(start_color='4A90D9', end_color='4A90D9', fill_type='solid')
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='DDE4ED'),
        right=Side(style='thin', color='DDE4ED'),
        top=Side(style='thin', color='DDE4ED'),
        bottom=Side(style='thin', color='DDE4ED'),
    )

    headers = ['序号', '作者', '日期', '原文位置', '批注内容']
    widths = [6, 16, 18, 50, 50]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = thin_border
        ws.column_dimensions[chr(64 + i)].width = w

    body_font = Font(name='微软雅黑', size=10)
    body_align = Alignment(vertical='center', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='center')
    even_fill = PatternFill(start_color='FAFBFC', end_color='FAFBFC', fill_type='solid')

    for row_idx, r in enumerate(results, 2):
        ws.row_dimensions[row_idx].height = 45
        values = [row_idx - 1, r['author'], r['date'], r['context'], r['text']]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = body_font
            cell.alignment = center_align if col_idx == 1 else body_align
            cell.border = thin_border
            if (row_idx - 2) % 2 == 1:
                cell.fill = even_fill

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:E{len(results) + 1}'
    wb.save(path)
    print(f"XLSX output: {path} ({len(results)} comments)")


def main():
    import argparse
    parser = argparse.ArgumentParser(description='Extract Word comments from DOCX files.')
    parser.add_argument('input', help='Path to the input DOCX file.')
    parser.add_argument('--output', choices=['json', 'xlsx'], default='xlsx',
                        help='Output format (default: xlsx).')
    parser.add_argument('--output-path', '-o', help='Output file path.')
    args = parser.parse_args()

    results = extract_comments(args.input)
    if not results:
        print("No comments found in the document.")
        return

    out_path = args.output_path
    if not out_path:
        base = os.path.splitext(os.path.basename(args.input))[0]
        ext = '.xlsx' if args.output == 'xlsx' else '.json'
        out_path = os.path.join(os.path.dirname(args.input), f'{base}_批注汇总{ext}')

    if args.output == 'xlsx':
        output_xlsx(results, out_path)
    else:
        output_json(results, out_path)


if __name__ == '__main__':
    main()
