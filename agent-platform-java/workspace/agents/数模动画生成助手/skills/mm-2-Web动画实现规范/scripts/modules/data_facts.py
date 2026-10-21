# -*- coding: utf-8 -*-
"""附件数据读取、事实摘要与图表数据解析。

这一层的目标不是替题目求解，而是把附件真正变成可被生成器引用的事实：
- 窄表（例如附件 1）提供完整序列与统计量；
- 宽表（例如「365 天 × 144 个时段」）提供按时刻平均、按天平均以及逐列序列；
- 所有事实都带稳定的 `facts:...` 引用键，供 HTML 图表直接使用。
"""
from __future__ import annotations

import csv
import io
import math
import re
import statistics
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


# ---------------------------------------------------------------------------
# 基础类型转换
# ---------------------------------------------------------------------------


def _clean_label(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    text = str(value).strip()
    return text or fallback


def _to_float(value: Any) -> Optional[float]:
    """把 Excel 单元格安全地转成数值；日期/时间/布尔值不当成数值。"""
    if value is None or isinstance(value, (bool,)):
        return None
    # datetime/date/time 都不应该被误识别为浮点数
    if hasattr(value, "hour") and not isinstance(value, (int, float)):
        return None
    if isinstance(value, (int, float)):
        value = float(value)
        return value if math.isfinite(value) else None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100.0
        except ValueError:
            return None
    try:
        value = float(text)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def _numbers(values: Iterable[Any]) -> List[float]:
    out: List[float] = []
    for value in values:
        number = _to_float(value)
        if number is not None:
            out.append(number)
    return out


def _stats(values: Iterable[Any]) -> Optional[Dict[str, Any]]:
    nums = _numbers(values)
    if not nums:
        return None
    return {
        "min": float(min(nums)),
        "max": float(max(nums)),
        "mean": float(statistics.fmean(nums)),
        "std": float(statistics.pstdev(nums)) if len(nums) > 1 else 0.0,
        "n": len(nums),
    }


def _mean_columns(rows: List[List[Optional[float]]]) -> List[float]:
    if not rows:
        return []
    width = max((len(row) for row in rows), default=0)
    out: List[float] = []
    for col in range(width):
        vals = [row[col] for row in rows if col < len(row) and row[col] is not None]
        out.append(float(statistics.fmean(vals)) if vals else 0.0)
    return out


def _mean_rows(rows: List[List[Optional[float]]]) -> List[float]:
    out: List[float] = []
    for row in rows:
        vals = [v for v in row if v is not None]
        out.append(float(statistics.fmean(vals)) if vals else 0.0)
    return out


# ---------------------------------------------------------------------------
# 表格读取
# ---------------------------------------------------------------------------


def _unique_headers(values: Iterable[Any]) -> List[str]:
    seen: Dict[str, int] = {}
    result: List[str] = []
    for index, value in enumerate(values):
        base = _clean_label(value, f"列{index + 1}")
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def _rows_to_table(name: str, rows: List[Any]) -> Optional[Dict[str, Any]]:
    """把二维单元格转成兼容旧接口、同时保留宽表结构的字典。"""
    if not rows:
        return None
    header = _unique_headers(list(rows[0] or []))
    if not header:
        return None
    data_rows = [list(row or []) for row in rows[1:]]
    ncol = len(header)

    # 每列数值（窄表和宽表都保留，供统计与单列图表使用）
    columns: Dict[str, List[float]] = {}
    for col_index, label in enumerate(header):
        vals = []
        for row in data_rows:
            vals.append(row[col_index] if col_index < len(row) else None)
        nums = _numbers(vals)
        if nums:
            columns[label] = nums

    # 找出「连续数值矩阵」：宽表的第一列通常是日期，前若干列可能是时间/发布时刻。
    counts: List[int] = []
    for col_index in range(ncol):
        counts.append(sum(1 for row in data_rows if col_index < len(row) and _to_float(row[col_index]) is not None))
    threshold = max(3, int(len(data_rows) * 0.60))
    matrix_indices = [i for i, count in enumerate(counts) if count >= threshold]
    matrix: List[List[Optional[float]]] = []
    row_labels: List[str] = []
    matrix_headers: List[str] = [header[i] for i in matrix_indices]
    if len(matrix_indices) >= 3:
        for row in data_rows:
            parsed = [_to_float(row[i]) if i < len(row) else None for i in matrix_indices]
            valid = sum(v is not None for v in parsed)
            # 允许少量空格/空单元格，但过滤完全空行
            if valid >= max(2, int(len(matrix_indices) * 0.55)):
                matrix.append(parsed)
                label = ""
                for i, value in enumerate(row):
                    if i in matrix_indices:
                        continue
                    if value not in (None, ""):
                        label = str(value).strip()
                        break
                row_labels.append(label)

    # 没有数值列的结果模板也需要被目录记录，因此返回一个空数值表，而不是 None。
    if not columns and not matrix:
        return {
            "name": name,
            "header": header,
            "nrow": len(data_rows),
            "columns": {},
            "matrix": [],
            "matrix_headers": [],
            "row_labels": [],
        }

    return {
        "name": name,
        "header": header,
        "nrow": len(data_rows),
        "columns": columns,
        "matrix": matrix,
        "matrix_headers": matrix_headers,
        "row_labels": row_labels,
    }


def read_xlsx_tables(path: Path) -> List[Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行 pip install -r requirements.txt") from exc

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    tables: List[Dict[str, Any]] = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            table = _rows_to_table(sheet_name, rows)
            if table is not None:
                tables.append(table)
    finally:
        wb.close()
    return tables


def read_csv_table(path: Path) -> List[Dict[str, Any]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    table = _rows_to_table(path.stem, rows)
    return [table] if table is not None else []


def read_docx_tables(path: Path) -> List[Dict[str, Any]]:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，请先执行 pip install -r requirements.txt") from exc
    doc = docx.Document(str(path))
    tables: List[Dict[str, Any]] = []
    for index, table in enumerate(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parsed = _rows_to_table(f"表{index + 1}", rows)
        if parsed is not None:
            tables.append(parsed)
    return tables


# ---------------------------------------------------------------------------
# 事实构建
# ---------------------------------------------------------------------------


def _attachment_label(name: str, path: Path) -> str:
    stem = path.stem or Path(name).stem
    # 竞赛题附件通常已经带「附件N」；保留这个短名，prompt 更稳定。
    if re.match(r"^附件\d+", stem):
        return stem
    return stem or path.name


def _namespace(label: str, sheet_name: str, sheet_count: int) -> str:
    if sheet_count <= 1 or not sheet_name or sheet_name.lower() in {"sheet1", "sheet"}:
        return label
    return f"{label}/{sheet_name}"


def _add_series(facts: Dict[str, Any], key: str, values: Iterable[Any]) -> None:
    nums = _numbers(values)
    if nums:
        facts.setdefault("_series", {})[key] = nums


def _add_stat(facts: Dict[str, Any], key: str, values: Iterable[Any]) -> None:
    stat = _stats(values)
    if stat:
        facts.setdefault("_stats", {})[key] = stat
        facts.setdefault("_series", {})[f"{key}:min"] = stat["min"]
        facts.setdefault("_series", {})[f"{key}:max"] = stat["max"]
        facts.setdefault("_series", {})[f"{key}:mean"] = stat["mean"]
        facts.setdefault("_series", {})[f"{key}:std"] = stat["std"]


def _read_attachment_tables(path: Path) -> List[Dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx_tables(path)
    if suffix == ".csv":
        return read_csv_table(path)
    if suffix == ".docx":
        return read_docx_tables(path)
    return []


def _add_table_facts(facts: Dict[str, Any], namespace: str, table: Dict[str, Any]) -> None:
    columns = table.get("columns") or {}
    matrix = table.get("matrix") or []
    matrix_headers = table.get("matrix_headers") or []
    facts.setdefault("_table_meta", {})[namespace] = {
        "rows": int(table.get("nrow") or 0),
        "columns": len(matrix_headers) if matrix_headers else len(columns),
        "headers": list(table.get("header") or [])[:20],
        "matrix": bool(matrix),
    }

    # 所有数值列都有统计；窄表直接暴露完整序列。
    for col, values in columns.items():
        key = f"{namespace}/{col}"
        _add_stat(facts, key, values)
        if len(values) <= 600:
            _add_series(facts, f"{key}:series", values)

    if not matrix or len(matrix_headers) < 12:
        return

    # 宽表：按列平均（一天的典型 144 格）、按行平均（全年每日均值）。
    slot_mean = _mean_columns(matrix)
    daily_mean = _mean_rows(matrix)
    _add_series(facts, f"{namespace}:slot_mean", slot_mean)
    _add_series(facts, f"{namespace}:daily_mean", daily_mean)
    for col_index, header in enumerate(matrix_headers):
        values = [row[col_index] for row in matrix if col_index < len(row) and row[col_index] is not None]
        if not values:
            continue
        key = f"{namespace}/{header}"
        _add_stat(facts, key, values)
        # 365 点以内的单列日序列可供需要时作图；超长数据仍然限制摘要展示。
        if len(values) <= 600:
            _add_series(facts, f"{key}:series", values)


def _find_series(facts: Dict[str, Any], key: str) -> Optional[List[float]]:
    value = (facts.get("_series") or {}).get(key)
    return value if isinstance(value, list) else None


def _add_derived_facts(facts: Dict[str, Any]) -> None:
    series = facts.setdefault("_series", {})
    # 兼容附件 1 与附件 2/4 的 144 段平均谱比较。
    for col in ("小区负载", "光伏发电预测功率"):
        a = _find_series(facts, f"附件1/{col}:series")
        for att in ("附件2", "附件4"):
            b = _find_series(facts, f"{att}/{col}:slot_mean")
            if a and b and len(a) == len(b):
                mae = sum(abs(x - y) for x, y in zip(a, b)) / len(a)
                series[f"附件1_vs_{att}/{col}:mae"] = float(mae)
    # 价格附件的全年日均曲线，名称明确，便于问题 4 使用。
    p4 = _find_series(facts, "附件4:daily_mean")
    if p4:
        series["附件4/日均电价:series"] = p4


def compute_facts(problem: Dict[str, Any]) -> Dict[str, Any]:
    """从 scan_problem_dir 返回的题目字典构建可引用事实。"""
    facts: Dict[str, Any] = {
        "_notes": [],
        "_series": {},
        "_stats": {},
        "_table_meta": {},
        "_catalog": [],
    }
    attachments = problem.get("attachments") or []

    for att in attachments:
        raw_path = att.get("path")
        if not raw_path:
            facts["_notes"].append(f"{att.get('name', '未知附件')} 缺少本地路径")
            continue
        path = Path(raw_path)
        label = _attachment_label(str(att.get("name") or path.name), path)
        try:
            tables = _read_attachment_tables(path)
        except Exception as exc:  # noqa: BLE001
            facts["_notes"].append(f"{att.get('name', path.name)} 读取失败：{exc}")
            tables = []

        catalog_entry = {
            "name": str(att.get("name") or path.name),
            "path": str(path),
            "label": label,
            "suffix": path.suffix.lower(),
            "size": int(att.get("size") or (path.stat().st_size if path.exists() else 0)),
            "sheets": [],
        }
        for table in tables:
            catalog_entry["sheets"].append({
                "name": table.get("name", ""),
                "rows": int(table.get("nrow") or 0),
                "columns": len(table.get("matrix_headers") or table.get("columns") or {}),
                "numeric": bool(table.get("columns") or table.get("matrix")),
            })
        facts["_catalog"].append(catalog_entry)

        sheet_count = len(tables)
        for table in tables:
            ns = _namespace(label, str(table.get("name") or ""), sheet_count)
            _add_table_facts(facts, ns, table)

    _add_derived_facts(facts)
    return facts


# ---------------------------------------------------------------------------
# Prompt 摘要与图表引用解析
# ---------------------------------------------------------------------------


def _fmt(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.6g}"
    return str(value)


def summarize(facts: Dict[str, Any], max_series_points: int = 12) -> str:
    series = facts.get("_series") or {}
    stats = facts.get("_stats") or {}
    catalog = facts.get("_catalog") or []
    lines: List[str] = []
    if catalog:
        lines.append("### 已读取附件（路径已由程序校验）")
        for item in catalog:
            sheets = item.get("sheets") or []
            shape = ", ".join(
                f"{s.get('name', '')} {s.get('rows', 0)}×{s.get('columns', 0)}"
                for s in sheets
            ) or "无可解析数值表"
            lines.append(f"- `{item.get('label')}` ← `{item.get('name')}`：{shape}")
    if not series and not stats:
        lines.append("\n（无可用数值附件；图表只能使用题面中的短序列。）")
        return "\n".join(lines)

    lines.append("\n### 可引用的真实序列（chart 的 series.d 使用 {\"from\":\"facts:键\"}）")
    grouped: Dict[str, List[str]] = {}
    for key in sorted(series):
        # 单值统计键太多时不占用序列展示配额
        if key.endswith((":min", ":max", ":mean", ":std")):
            continue
        group = key.split(":", 1)[0].split("/", 1)[0]
        grouped.setdefault(group, []).append(key)
    for group, keys in grouped.items():
        lines.append(f"\n#### {group}")
        # 优先展示能直接画图的典型序列
        keys = sorted(keys, key=lambda k: (0 if "slot_mean" in k else 1 if "daily_mean" in k else 2, k))
        for key in keys[:10]:
            value = series[key]
            if isinstance(value, list):
                head = ", ".join(_fmt(v) for v in value[:max_series_points])
                lines.append(f"- `facts:{key}`：长度 {len(value)}；前 {min(len(value), max_series_points)} 项 {head}{' …' if len(value) > max_series_points else ''}")
            else:
                lines.append(f"- `facts:{key}` = {_fmt(value)}")

    lines.append("\n### 关键统计（可直接写进说明）")
    for key, stat in sorted(stats.items()):
        if ":" in key:
            continue
        lines.append(
            f"- `{key}`：min={_fmt(stat['min'])} max={_fmt(stat['max'])} "
            f"mean={_fmt(stat['mean'])} std={_fmt(stat['std'])} n={stat['n']}"
        )
        if len(lines) >= 80:
            break
    notes = facts.get("_notes") or []
    if notes:
        lines.append("\n### 数据读取备注")
        lines.extend(f"- {note}" for note in notes[:12])
    return "\n".join(lines)


def resolve_fact_ref(facts: Dict[str, Any], ref: str) -> Any:
    if not ref or not isinstance(ref, str) or not ref.startswith("facts:"):
        return None
    return (facts.get("_series") or {}).get(ref[len("facts:"):])


def resolve_chart_data(facts: Dict[str, Any], deck: Dict[str, Any]) -> None:
    """就地解析所有 {from:"facts:..."}，并记录未命中的引用。"""
    notes = facts.setdefault("_notes", [])

    def walk(value: Any) -> Any:
        if isinstance(value, dict):
            if isinstance(value.get("from"), str):
                resolved = resolve_fact_ref(facts, value["from"])
                if resolved is None:
                    notes.append(f"未解析图表数据引用：{value['from']}")
                    return []
                return resolved
            for key in list(value.keys()):
                value[key] = walk(value[key])
            return value
        if isinstance(value, list):
            return [walk(item) for item in value]
        return value

    walk(deck)
